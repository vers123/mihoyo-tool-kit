"""新闻 Excel 导出。

从 SQLite 读取四站点新闻，每站点一个 sheet，输出 ``output/news.xlsx``。
openpyxl 为可选依赖（``mihoyo-toolkit[excel]``），故在导出函数内部懒加载。
"""

from __future__ import annotations

from pathlib import Path

from ..core.config import get_settings
from ..core.paths import get_path_manager
from ..core.storage import Storage
from ..utils.logger import get_module_logger

logger = get_module_logger("exporters.excel")

#: (表头, 列宽)
_COLUMNS: list[tuple[str, int]] = [
    ("ID", 10),
    ("标题", 40),
    ("日期", 20),
    ("分类", 12),
    ("摘要", 50),
    ("封面", 30),
    ("链接", 50),
]


class ExcelWriter:
    """四站点新闻 → Excel 导出器。"""

    def __init__(
        self,
        output_path: str | Path | None = None,
        *,
        db_path: str | Path | None = None,
    ) -> None:
        path_manager = get_path_manager()
        self.output_path = Path(output_path) if output_path else path_manager.output / "news.xlsx"
        self._db_path = db_path

    def export(self) -> Path:
        """读取 SQLite 并写出 xlsx，返回输出路径。"""
        # openpyxl 懒加载：仅导出时导入，避免非 Excel 场景强依赖
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        settings = get_settings()
        path_manager = get_path_manager()
        self.output_path.parent.mkdir(parents=True, exist_ok=True)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_align = Alignment(horizontal="center", vertical="center")

        workbook = Workbook()
        default_sheet = workbook.active
        workbook.remove(default_sheet)

        total = 0
        with Storage(self._db_path) as store:
            for game in settings.sources.news.keys():  # noqa: SIM118 - 自定义方法，非 dict
                site = settings.sources.news.get_site(game)
                items = store.query_news(game)
                sheet = workbook.create_sheet(title=site.label)

                for col, (header, _) in enumerate(_COLUMNS, 1):
                    cell = sheet.cell(row=1, column=col, value=header)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_align

                for row_idx, item in enumerate(items, 2):
                    values = (
                        item.iInfoId,
                        item.sTitle,
                        item.dtStartTime,
                        item.sCategoryName,
                        item.sIntro,
                        item.poster_url,
                        item.url,
                    )
                    for col_idx, value in enumerate(values, 1):
                        sheet.cell(row=row_idx, column=col_idx, value=value)

                for col, (_, width) in enumerate(_COLUMNS, 1):
                    sheet.column_dimensions[get_column_letter(col)].width = width
                sheet.freeze_panes = "A2"
                total += len(items)

        workbook.save(self.output_path)
        logger.info(
            "Excel 已导出: %s（%d 个 sheet，共 %d 条）",
            path_manager.relative(self.output_path),
            len(settings.sources.news.keys()),
            total,
        )
        return self.output_path


def export_news_excel() -> Path:
    """导出四站点新闻到 ``output/news.xlsx``。"""
    return ExcelWriter().export()


__all__ = ["ExcelWriter", "export_news_excel"]
