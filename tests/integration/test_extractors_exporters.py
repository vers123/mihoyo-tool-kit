"""提取器与导出器集成测试（SQLite + 文件输出）。"""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from mihoyo_toolkit.core.models import ImageItem, PostItem, TutorialItem, WeiboItem
from mihoyo_toolkit.core.paths import get_path_manager
from mihoyo_toolkit.exporters import (
    ExcelWriter,
    FeedGenerator,
    export_news_excel,
    generate_json_feed,
    generate_rss_feed,
)
from mihoyo_toolkit.extractors import (
    ChangelogExtractor,
    GenshinENNewsExtractor,
    GenshinNewsExtractor,
    ImageExtractor,
    PostExtractor,
    SRNewsExtractor,
    TutorialExtractor,
    WeiboExtractor,
    ZZZNewsExtractor,
    run_extract_images,
    run_extract_news,
    run_extract_posts,
    run_extract_tutorial,
    run_extract_weibo,
)

TUTORIAL_HTML = (
    "<html><table>"
    '<tr class="table-row"><td><p>0</p></td><td><p><span>角色名</span></p></td></tr>'
    '<tr class="table-row"><td><p><span>10000001</span></p></td>'
    "<td><p><span>角色A</span></p></td></tr>"
    "</table></html>"
)

CHANGELOG_HTML = """
<html><body>
<h2>7.0版本-2026/08/12</h2>
<h3>一、内容新增</h3>
<p>新增说明</p>
<h4>1. 条目A</h4>
<p>条目A说明</p>
<a href="https://act.mihoyo.com/ys/ugc/tutorial/detail/abc123">条目A链接</a>
<h3>二、内容修改</h3>
<p>修改说明</p>
<a href="https://act.mihoyo.com/ys/ugc/tutorial/detail/def456">修改链接</a>
<h2>6.0版本-2026/01/01</h2>
<h3>内容新增</h3>
<p>old</p>
</body></html>
"""


# ---------------------------------------------------------------------- #
#  新闻提取
# ---------------------------------------------------------------------- #
def test_news_extractor_export(store, sample_news_items, tmp_home: Path) -> None:
    store.upsert_news("genshin", sample_news_items)

    extractor = GenshinNewsExtractor()
    assert extractor.game == "genshin"
    assert extractor.output_path == tmp_home / "data" / "results" / "genshin_news.txt"

    items = extractor.extract_news()
    assert len(items) == 2
    assert len(extractor.load_news(limit=1)) == 1

    line = extractor.format_line(1, items[0])
    assert line.startswith("0001-")
    assert "https://ys.mihoyo.com/main/news/detail/" in line

    out = extractor.export()
    assert out is not None
    assert out.is_file()
    assert "版本更新说明" in out.read_text(encoding="utf-8")


def test_news_extractor_no_data_returns_none(store) -> None:
    assert ZZZNewsExtractor().export() is None


def test_news_extractor_subclasses() -> None:
    assert GenshinENNewsExtractor().game == "genshin_en"
    assert SRNewsExtractor().game == "starrail"
    assert ZZZNewsExtractor().game == "zzz"


def test_run_extract_news(store, sample_news_items, tmp_home: Path) -> None:
    store.upsert_news("genshin", sample_news_items)
    run_extract_news("genshin")
    assert (tmp_home / "data" / "results" / "genshin_news.txt").is_file()

    run_extract_news("unknown")  # 未知站点：记录错误后返回


# ---------------------------------------------------------------------- #
#  发帖 / 微博 / 图片提取
# ---------------------------------------------------------------------- #
def test_post_extractor(store, tmp_home: Path) -> None:
    store.upsert_posts(
        [
            PostItem(post_id="p1", title="标题1", created_at="2024-01-02", url="u1"),
            PostItem(post_id="p2", title="标题2", created_at="2024-01-01", url="u2"),
        ]
    )
    extractor = PostExtractor()
    items = extractor.extract_posts(limit=1)
    assert len(items) == 1
    assert items[0].post_id == "p1"

    out = extractor.save_post_data(items)
    assert out == tmp_home / "data" / "results" / "posts.txt"

    run_extract_posts()


def test_post_extractor_empty(store) -> None:
    assert PostExtractor().extract_posts() == []
    run_extract_posts()  # 无数据：记录日志后返回


def test_weibo_extractor(store, tmp_home: Path) -> None:
    store.upsert_weibo([WeiboItem(post_id="w1", text="正文", created_at="2024-01-01", url="u1")])
    extractor = WeiboExtractor()
    items = extractor.extract_weibo(limit=1)
    assert len(items) == 1

    out = extractor.save_weibo_data(items)
    assert out == tmp_home / "data" / "results" / "weibo.txt"
    run_extract_weibo()


def test_weibo_extractor_empty(store) -> None:
    assert WeiboExtractor().extract_weibo() == []
    run_extract_weibo()


def test_image_extractor(store, tmp_home: Path) -> None:
    store.upsert_images([ImageItem(character_id="c1", name="角色A", image_url="https://x/1.png")])
    extractor = ImageExtractor()
    items = extractor.extract_image_urls(limit=1)
    assert len(items) == 1

    out = extractor.save_image_data(items)
    assert out == tmp_home / "data" / "results" / "image_urls.txt"
    run_extract_images()


def test_image_extractor_empty(store) -> None:
    assert ImageExtractor().extract_image_urls() == []
    run_extract_images()


# ---------------------------------------------------------------------- #
#  教程提取
# ---------------------------------------------------------------------- #
def test_tutorial_extractor_characters(store, tmp_home: Path) -> None:
    html_dir = get_path_manager().html / "tutorial"
    html_dir.mkdir(parents=True, exist_ok=True)
    (html_dir / "tutorial_tid1.html").write_text(TUTORIAL_HTML, encoding="utf-8")

    extractor = TutorialExtractor("tid1")
    assert extractor.html_path.is_file()
    assert extractor.read_html() == TUTORIAL_HTML

    items = extractor.extract_characters()
    assert items == [TutorialItem(character_id="10000001", name="角色A", lang="")]

    assert extractor.store(items) == 1
    out = extractor.save_character_data(items)
    assert out == tmp_home / "data" / "results" / "characters_tid1.txt"
    assert out.is_file()


def test_tutorial_extractor_loose_pattern(store) -> None:
    extractor = TutorialExtractor("tid2")
    html = "<td><p>10000002</p></td><td><p>角色B</p></td>"
    items = extractor.extract_characters(html)
    assert items == [TutorialItem(character_id="10000002", name="角色B", lang="")]


def test_tutorial_extractor_missing_html(tmp_home: Path) -> None:
    from mihoyo_toolkit.core.exceptions import ParseError

    with pytest.raises(ParseError):
        TutorialExtractor("missing").read_html()


def test_run_extract_tutorial_characters(store, tmp_home: Path) -> None:
    html_dir = get_path_manager().html / "tutorial"
    html_dir.mkdir(parents=True, exist_ok=True)
    (html_dir / "tutorial_run1.html").write_text(TUTORIAL_HTML, encoding="utf-8")

    run_extract_tutorial("run1")
    assert store.count_table("tutorial") == 1


def test_run_extract_tutorial_missing(tmp_home: Path) -> None:
    run_extract_tutorial("does-not-exist")  # 缺失：记录错误后返回


def test_run_extract_tutorial_no_characters(store, tmp_home: Path) -> None:
    html_dir = get_path_manager().html / "tutorial"
    html_dir.mkdir(parents=True, exist_ok=True)
    (html_dir / "tutorial_empty1.html").write_text("<html>无数据</html>", encoding="utf-8")
    run_extract_tutorial("empty1")


def test_changelog_extractor(tmp_home: Path) -> None:
    assert ChangelogExtractor.is_changelog(CHANGELOG_HTML) is True
    assert ChangelogExtractor.is_changelog("普通页面") is False

    extractor = ChangelogExtractor("tid3")
    links = extractor.extract_all_links(CHANGELOG_HTML)
    assert {link["tutorial_id"] for link in links} == {"abc123", "def456"}

    data = extractor.extract(CHANGELOG_HTML)
    assert data["page_id"] == "tid3"
    assert len(data["versions"]) == 2
    assert data["versions"][0]["version"] == "7.0版本"
    assert data["versions"][0]["date"] == "2026/08/12"
    assert len(data["versions"][0]["categories"]) == 2

    out = extractor.save(data)
    assert out.is_file()
    assert json.loads(out.read_text(encoding="utf-8"))["page_id"] == "tid3"


def test_run_extract_tutorial_changelog(tmp_home: Path) -> None:
    html_dir = get_path_manager().html / "tutorial"
    html_dir.mkdir(parents=True, exist_ok=True)
    (html_dir / "tutorial_cl1.html").write_text(CHANGELOG_HTML, encoding="utf-8")

    run_extract_tutorial("cl1")
    assert (get_path_manager().results / "changelog_cl1.json").is_file()


# ---------------------------------------------------------------------- #
#  导出器
# ---------------------------------------------------------------------- #
def test_export_news_excel(store, sample_news_items, tmp_home: Path) -> None:
    store.upsert_news("genshin", sample_news_items)

    writer = ExcelWriter()
    assert writer.output_path == tmp_home / "output" / "news.xlsx"
    out = writer.export()
    assert out.is_file()

    from openpyxl import load_workbook

    workbook = load_workbook(out)
    assert "原神" in workbook.sheetnames
    sheet = workbook["原神"]
    assert sheet["A1"].value == "ID"
    assert sheet["A2"].value == 1001


def test_export_news_excel_wrapper(store, sample_news_items) -> None:
    store.upsert_news("genshin", sample_news_items)
    assert export_news_excel().is_file()


def test_feed_generator(store, sample_news_items, tmp_home: Path) -> None:
    store.upsert_news("genshin", sample_news_items)

    generator = FeedGenerator()
    rss_path = generator.write_rss()
    json_path = generator.write_json()

    assert rss_path == tmp_home / "output" / "news_feed.xml"
    assert json_path == tmp_home / "output" / "news_feed.json"

    rss_text = rss_path.read_text(encoding="utf-8")
    assert "<rss" in rss_text
    assert "版本更新说明" in rss_text
    assert "genshin-1001" in rss_text

    payload = json.loads(json_path.read_text(encoding="utf-8"))
    assert payload["version"] == "https://jsonfeed.org/version/1.2"
    assert len(payload["items"]) == 2
    assert payload["items"][0]["id"].startswith("genshin-")


def test_feed_wrappers(store, sample_news_items) -> None:
    store.upsert_news("genshin", sample_news_items)
    assert generate_rss_feed().is_file()
    assert generate_json_feed().is_file()
